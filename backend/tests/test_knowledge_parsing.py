import io
import zipfile
import pytest
from app.config import Settings
from app.errors import AppError
from app.knowledge.config import KnowledgeSettings
from tests.test_material_parsers import pdf_bytes


def parse(name,data,**kwargs):
    from app.knowledge.parsing import parse_file
    return parse_file(name,data,Settings(_env_file=None,**kwargs))


def test_pdf_keeps_actual_page_even_after_empty_page_and_docx_keeps_block_order():
    result = parse('pages.pdf',pdf_bytes(['First fact','','Third fact']))
    assert [s['page'] for s in result['segments']] == [1,3]
    from docx import Document
    document = Document()
    document.add_paragraph('第一段')
    document.add_table(1,1).cell(0,0).text = '中间表格'
    document.add_paragraph('末段')
    data = io.BytesIO()
    document.save(data)
    result = parse('note.docx',data.getvalue())
    assert [s['text'] for s in result['segments']] == ['第一段','中间表格','末段']
    assert all('page' not in s for s in result['segments'])
    for segment in result['segments']:
        assert result['text'][segment['start_index']:segment['end_index']] == segment['text']


@pytest.mark.parametrize('encoding',['utf-8-sig','gb18030'])
def test_markdown_keeps_heading_code_table_html_as_plain_text(encoding):
    text = '# 标题\n\n中文说明\n\n```python\n# 代码中的注释\nprint(1)\n```\n\n|列|值|\n|--|--|\n|a|b|\n<script>alert(1)</script>'
    result = parse('note.md',text.encode(encoding))
    assert result['text'] == text
    assert result['source_type'] == 'md' and 'print(1)' in result['text']
    assert result['parser_version'] == 'knowledge-parser-v1'
    assert parse('note.txt','中文\n第二行'.encode(encoding))['text'] == '中文\n第二行'


def test_corrupt_empty_encrypted_scan_and_expansion_limits():
    from pypdf import PdfWriter
    blank = PdfWriter()
    blank.add_blank_page(100,100)
    stream = io.BytesIO()
    blank.write(stream)
    for name,data in [('scan.pdf',stream.getvalue()),('bad.pdf',b'%PDF-invalid'),('a.md',b'  '),('binary.md',b'\x00')]:
        with pytest.raises(AppError): parse(name,data)
    blank.encrypt('password')
    stream = io.BytesIO()
    blank.write(stream)
    with pytest.raises(AppError): parse('locked.pdf',stream.getvalue())
    with pytest.raises(AppError): parse('large.md',b'x'*11,material_max_chars=10)
    with pytest.raises(AppError): parse('many.pdf',pdf_bytes(['a','b']),material_max_pages=1)
    stream = io.BytesIO()
    with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml',b'x'*(50*1024*1024+1))
    with pytest.raises(AppError): parse('bomb.docx',stream.getvalue())


def test_recursive_chunks_are_exact_stable_bounded_and_never_cross_pages_or_headers():
    from app.knowledge.parsing import split_content
    config = KnowledgeSettings(_env_file=None,chunk_size=100,chunk_overlap=20)
    text = '# 第一章\n\n'+'中文长段落内容，解释概念。'*40+'\n\n## 第二节\n\n'+'新的主题。'*80
    parsed = parse('note.md',text.encode())
    first = split_content(parsed,'revision-one',config)
    assert first == split_content(parsed,'revision-one',config)
    assert len(first)>5 and all(0 < len(c['text']) <= 100 for c in first)
    for chunk in first:
        location = chunk['location']
        assert parsed['text'][location['start_index']:location['end_index']] == chunk['text']
        assert len(chunk['id']) == len(chunk['content_hash']) == 64
        assert location['header_path'][0] == '第一章'
        assert not ('新的主题' in chunk['text'] and '中文长段落' in chunk['text'])
    assert any(a['location']['end_index'] > b['location']['start_index'] for a,b in zip(first,first[1:]))
    other = split_content(parsed,'revision-two',config)
    assert first[0]['id'] != other[0]['id'] and first[0]['content_hash'] == other[0]['content_hash']
    pdf = split_content(parse('x.pdf',pdf_bytes(['A '*100,'B '*100])),'pdf-revision',config)
    assert {c['location']['page'] for c in pdf} == {1,2}
    assert all(not ('A' in c['text'] and 'B' in c['text']) for c in pdf)


def test_repeated_text_and_fenced_headings_have_correct_original_locations():
    from app.knowledge.parsing import split_content
    text = '# 标题\n\n重复内容。\n\n```md\n# 这不是章节\n```\n\n重复内容。\n\n# 另一个标题\n\n重复内容。'
    result = parse('a.markdown',text.encode())
    chunks = split_content(result,'r',KnowledgeSettings(_env_file=None,chunk_size=100,chunk_overlap=10))
    assert all(c['location']['header_path'] in (['标题'],['另一个标题']) for c in chunks)
    assert all(text[c['location']['start_index']:c['location']['end_index']] == c['text'] for c in chunks)
