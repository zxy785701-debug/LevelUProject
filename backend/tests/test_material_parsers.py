import io
import zipfile
import pytest
from app.errors import AppError
from app.materials.parsers import parse_document, validate_url, caption_segments


def pdf_bytes(texts):
    # Real text PDF, without requiring a test-only rendering service.
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=300, height=400)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 12 Tf 30 300 Td ({text}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output)
    return output.getvalue()


def test_pdf_ranges_preserve_page_references_and_no_silent_truncation():
    result = parse_document('notes.pdf', pdf_bytes(['First fact', 'Second fact', 'Third fact']), pages='2-3')
    assert [s['page'] for s in result['segments']] == [2, 3]
    assert 'Second fact' in result['segments'][0]['text']
    assert 'First fact' not in str(result)
    for pages in ['0', '3-2', '1,9', 'x', '1-2,2']:
        with pytest.raises(AppError): parse_document('a.pdf', pdf_bytes(['One', 'Two']), pages=pages)
    with pytest.raises(AppError, match='文字'):
        from pypdf import PdfWriter
        blank = io.BytesIO(); writer = PdfWriter(); writer.add_blank_page(300, 400); writer.write(blank)
        parse_document('scan.pdf', blank.getvalue())
    with pytest.raises(AppError): parse_document('a.pdf', pdf_bytes(['Too much text']), max_chars=2)


def test_docx_paragraph_table_order_and_txt_encoding():
    from docx import Document
    doc = Document(); doc.add_paragraph('第一段'); doc.add_table(1, 1).cell(0, 0).text = '表格内容'; doc.add_paragraph('末段')
    stream = io.BytesIO(); doc.save(stream)
    result = parse_document('a.docx', stream.getvalue())
    text = '\n'.join(s['text'] for s in result['segments'])
    assert text.index('第一段') < text.index('表格内容') < text.index('末段')
    assert all(s.get('page') is None for s in result['segments'])
    for encoding in ('utf-8-sig', 'gb18030'):
        assert '中文笔记' in parse_document('a.txt', '中文笔记'.encode(encoding))['segments'][0]['text']
    with pytest.raises(AppError): parse_document('a.docx', stream.getvalue(), pages='1')


@pytest.mark.parametrize('name,data', [('a.doc', b'old'), ('fake.pdf', b'not pdf'), ('fake.docx', b'not zip'), ('empty.txt', b''), ('binary.txt', b'\x00\x01')])
def test_reject_unsupported_corrupt_empty_and_binary_files(name, data):
    with pytest.raises(AppError): parse_document(name, data)


def test_zip_expansion_and_file_limit():
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', b'x' * 2000)
    with pytest.raises(AppError): parse_document('bomb.docx', stream.getvalue(), max_expanded=1000)
    with pytest.raises(AppError): parse_document('a.txt', b'x' * 20, max_bytes=10)


@pytest.mark.parametrize('url', ['file:///etc/passwd', 'http://127.0.0.1', 'http://[::1]', 'http://169.254.169.254/latest', 'https://user:pass@example.com', 'http://2130706433', 'http://example.com:3306', 'https://localhost', 'http://10.1.2.3'])
def test_url_blocks_local_and_unsupported_targets(url):
    with pytest.raises(AppError): validate_url(url, resolve=lambda host: ['127.0.0.1'])


def test_dns_mixed_addresses_blocked_and_public_url_canonicalized():
    with pytest.raises(AppError): validate_url('https://example.com', resolve=lambda host: ['93.184.216.34', '::1'])
    assert validate_url('https://example.com/article#part', resolve=lambda host: ['93.184.216.34']) == 'https://example.com/article'


def test_caption_range_deduplication_and_timestamps():
    vtt = 'WEBVTT\n\n00:00:01.000 --> 00:00:03.000\nFirst\n\n00:00:04.000 --> 00:00:06.000\nSecond &amp; third\n'
    result = caption_segments(vtt, 'vtt', start=3, end=6)
    assert len(result) == 1 and result[0]['start'] == 4
    assert result[0]['text'] == 'Second & third'
    data = '{"events":[{"tStartMs":1000,"dDurationMs":1000,"segs":[{"utf8":"你好"}]}]}'
    assert caption_segments(data, 'json3', 0, 3)[0]['text'] == '你好'
    with pytest.raises(AppError): caption_segments(vtt, 'vtt', 8, 9)
